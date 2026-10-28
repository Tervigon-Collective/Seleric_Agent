"""Excel datasets are built from the evidence the loop stored (the original rows), never from display strings.

The fixtures go through the real writers (``write_evidence`` / ``_persist_turn``), so a change to what is stored breaks
these tests instead of silently emptying the attachment."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from seleric_swarm.agent.dependencies import ExecutionLimits, NullMcpClient, SelericDeps
from seleric_swarm.agent.mcp_toolset import write_evidence, write_query_evidence
from seleric_swarm.agent.runner import _persist_turn
from seleric_swarm.conversations.contracts import ContextBundle, Principal
from seleric_swarm.mail.evidence import datasets_for
from seleric_swarm.mail.excel import build_workbook
from seleric_swarm.services.catalogue_bootstrap import CatalogueMetricMeta, CatalogueSnapshot
from seleric_swarm.state.artifacts import InMemoryArtifactStore

CATALOGUE = CatalogueSnapshot(metrics=(
    CatalogueMetricMeta(id="net_sales", label="Net sales", raw={"unit": "INR", "description": "Net sales after returns"}),
    CatalogueMetricMeta(id="orders", label="Orders", raw={"unit": "count"}),
))
WINDOW = {"time_range": {"start": "2026-10-21", "end": "2026-10-27"}, "metric_ids": ["net_sales"], "currency": "INR"}
CHANNELS = {
    "query_id": "q_ch", "columns": ["channel", "ad_platform", "net_sales"],
    "rows": [{"channel": "ig_feed", "ad_platform": "meta", "net_sales": "131269.36"},
             {"channel": "google_pmax", "ad_platform": "google", "net_sales": "103520.57"},
             {"channel": "fb_feed", "ad_platform": "meta", "net_sales": "37977.99"}],
    "provenance": {**WINDOW, "query_id": "q_ch"},
}


def _deps(store: InMemoryArtifactStore, mission: str, thread: str = "t1") -> SelericDeps:
    return SelericDeps(
        mission_id=mission, as_of=datetime(2026, 10, 28, tzinfo=ZoneInfo("Asia/Kolkata")),
        principal=Principal(principal_id="p1", workspace_id="ws1", user_id="u1"), thread_id=thread, run_id="r1",
        trace_id="tr", context=ContextBundle(), mcp_client=NullMcpClient(), artifact_store=store,
        limits=ExecutionLimits(), catalogue=CATALOGUE)


def _finish_turn(store, deps, evidence_ids) -> None:
    state = SimpleNamespace(evidence=[], evidence_ids=evidence_ids, payloads=[])
    _persist_turn(store, loop_result=SimpleNamespace(state=state), query="q", workspace_id="ws1",
                  thread_id=deps.thread_id, mission_id=deps.mission_id, as_of_dt=deps.as_of)


def _sheet(datasets, name_part: str):
    wb = load_workbook(BytesIO(build_workbook(datasets)))
    return wb[next(n for n in wb.sheetnames if name_part in n)], wb["About"]


def test_a_stored_query_becomes_a_typed_sheet_with_the_original_numbers() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    write_evidence(CHANNELS, {"measures": ["net_sales"]}, deps)
    datasets, missing = datasets_for(deps, ["q_ch"])
    assert missing == [] and len(datasets) == 1
    ws, about = _sheet(datasets, "Net sales")
    header = [c.value for c in ws[1]]
    assert "Channel" in header and "Ad Platform" in header and any(h.startswith("Net sales") for h in header)
    values = {ws.cell(r, header.index("Channel") + 1).value: ws.cell(r, next(i for i, h in enumerate(header, 1) if h.startswith("Net sales"))).value
              for r in range(2, ws.max_row + 1)}
    assert values == {"ig_feed": 131269.36, "google_pmax": 103520.57, "fb_feed": 37977.99}  # numbers, not "₹1,31,269"
    text = "\n".join(str(c.value) for row in about.iter_rows() for c in row if c.value)
    assert "2026-10-21 to 2026-10-27" in text and "Asia/Kolkata" in text and "Net sales after returns" in text


def test_a_query_from_an_earlier_turn_is_found_through_the_threads_turn_records() -> None:
    store = InMemoryArtifactStore()
    first = _deps(store, "m1")
    ids = write_evidence(CHANNELS, {"measures": ["net_sales"]}, first)
    _finish_turn(store, first, ids)
    datasets, missing = datasets_for(_deps(store, "m2"), ["q_ch"])  # a later turn of the same thread
    assert missing == [] and len(datasets) == 1
    assert datasets_for(_deps(store, "m3", thread="other-thread"), ["q_ch"]) == ([], ["q_ch"])  # never another thread's


def test_this_turns_parent_query_id_expands_to_all_its_parts() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    part_b = {**CHANNELS, "query_id": "q_b", "columns": ["channel", "orders"],
              "rows": [{"channel": "ig_feed", "orders": "40"}],
              "provenance": {**WINDOW, "query_id": "q_b", "metric_ids": ["orders"]}}
    parent = {"query_id": "q_parent", "parts": [CHANNELS, part_b]}
    by_query = write_query_evidence(parent, {"measures": ["net_sales", "orders"]}, deps)  # what the turn's state keeps
    datasets, missing = datasets_for(deps, ["q_parent"], by_query=by_query)  # no turn_record yet: this is the running turn
    assert missing == [] and len(datasets) == 2


def test_an_earlier_turns_parent_query_id_finds_its_parts_through_the_turn_payloads() -> None:
    store = InMemoryArtifactStore()
    first = _deps(store, "m1")
    parent = {"query_id": "q_parent", "parts": [CHANNELS]}
    ids = write_query_evidence(parent, {"measures": ["net_sales"]}, first)["q_ch"]
    state = SimpleNamespace(evidence=[], evidence_ids=ids, payloads=[parent])
    _persist_turn(store, loop_result=SimpleNamespace(state=state), query="q", workspace_id="ws1",
                  thread_id="t1", mission_id="m1", as_of_dt=first.as_of)
    datasets, missing = datasets_for(_deps(store, "m2"), ["q_parent"])  # a later turn: no by_query map
    assert missing == [] and len(datasets) == 1


def test_a_cut_off_query_is_marked_partial_in_the_workbook() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    write_evidence(CHANNELS, {"measures": ["net_sales"]}, deps, max_points=2)
    datasets, _ = datasets_for(deps, ["q_ch"])
    assert datasets[0].total_rows == 3 and len(datasets[0].rows) == 2
    _, about = _sheet(datasets, "Net sales")
    text = "\n".join(str(c.value) for row in about.iter_rows() for c in row if c.value)
    assert "PARTIAL" in text and "2 of 3" in text


def test_a_query_with_no_stored_rows_is_reported_missing_not_dropped() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    write_evidence(CHANNELS, {"measures": ["net_sales"]}, deps)
    datasets, missing = datasets_for(deps, ["q_ch", "q_python_result"])
    assert len(datasets) == 1 and missing == ["q_python_result"]


def test_compare_windows_get_a_period_column_and_share_ratios_are_percentages() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    payload = {"query_id": "q_cmp", "columns": ["net_sales"], "rows": [{"net_sales": "342703.15"}],
               "compare_rows": [{"net_sales": "384242.57"}],
               "provenance": {**WINDOW, "query_id": "q_cmp", "compare_period": {"start": "2026-10-14", "end": "2026-10-20"}}}
    write_evidence(payload, {"measures": ["net_sales"]}, deps)
    (ds,), _ = datasets_for(deps, ["q_cmp"])
    assert [c.header for c in ds.columns] == ["Period start", "Period end", "Period", "Net sales (INR)"]
    assert {(r["period"], r["net_sales"]) for r in ds.rows} == {("current", 342703.15), ("previous", 384242.57)}


# ---- tables in the email body: the same stored rows, formatted the way answers format them

from seleric_swarm.mail.evidence import tables_for


def test_a_stored_query_becomes_a_body_table_with_display_formatting() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    write_evidence(CHANNELS, {"measures": ["net_sales"]}, deps)
    tables, missing = tables_for(deps, ["q_ch"])
    (t,) = tables["q_ch"]
    assert missing == [] and t.title == "Net sales by Channel and Ad Platform" and t.subtitle == "2026-10-21 to 2026-10-27"
    assert t.headers == ["Channel", "Ad Platform", "Net sales (INR)"] and t.align == ["l", "l", "r"]
    assert t.rows[0] == ["ig_feed", "meta", "₹1,31,269"] and t.rows[2][2] == "₹37,978"  # the period is in the subtitle, not a column


def test_a_long_table_is_cut_to_max_rows_and_says_so() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    write_evidence(CHANNELS, {"measures": ["net_sales"]}, deps)
    (t,) = tables_for(deps, ["q_ch"], max_rows=2)[0]["q_ch"]
    assert len(t.rows) == 2 and t.note == "Showing 2 of 3 rows."
    (t,) = tables_for(deps, ["q_ch"], max_rows=2, attached=True)[0]["q_ch"]
    assert t.note == "Showing 2 of 3 rows. The full data is in the attachment."


def test_stored_evidence_that_was_cut_is_called_partial_in_the_table() -> None:
    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    write_evidence(CHANNELS, {"measures": ["net_sales"]}, deps, max_points=2)
    (t,) = tables_for(deps, ["q_ch"])[0]["q_ch"]
    assert "Partial data" in t.note and "2 of 3" in t.note


def test_shares_counts_and_missing_values_read_like_the_answers() -> None:
    store = InMemoryArtifactStore()
    catalogue = CatalogueSnapshot(metrics=(CATALOGUE.metrics[0], CatalogueMetricMeta(id="cvr", label="Conversion rate", raw={"unit": "percent"}),
                                           CatalogueMetricMeta(id="orders", label="Orders", raw={"unit": "count"})))
    deps = _deps(store, "m1")
    object.__setattr__(deps, "catalogue", catalogue)
    payload = {"query_id": "q_m", "columns": ["channel", "cvr", "orders"],
               "rows": [{"channel": "a", "cvr": "3.14", "orders": "1200"}, {"channel": "b", "cvr": "1", "orders": None}],
               "provenance": {**WINDOW, "query_id": "q_m", "metric_ids": ["cvr", "orders"]}}
    write_evidence(payload, {"measures": ["cvr", "orders"]}, deps)
    (t,) = tables_for(deps, ["q_m"])[0]["q_m"]
    cells = {r[0]: r[1:] for r in t.rows}
    assert cells["a"] == ["3.1%", "1,200"] and cells["b"][0] == "1.0%" and cells["b"][1] in ("—", "-")  # no data is not 0


def test_a_table_for_a_query_with_nothing_stored_is_reported_missing() -> None:
    deps = _deps(InMemoryArtifactStore(), "m1")
    assert tables_for(deps, ["q_nope"]) == ({}, ["q_nope"])


def test_a_daily_series_becomes_an_inline_chart_and_a_breakdown_does_not() -> None:
    from seleric_swarm.mail.evidence import charts_for

    store = InMemoryArtifactStore()
    deps = _deps(store, "m1")
    daily = {"query_id": "q_day", "columns": ["orders.created_at.day", "net_sales"],
             "rows": [{"orders.created_at.day": f"2026-10-2{d}", "net_sales": str(1000 * d)} for d in range(1, 8)],
             "provenance": {**WINDOW, "query_id": "q_day"}}
    write_evidence(daily, {"measures": ["net_sales"], "grain": "day"}, deps)
    write_evidence(CHANNELS, {"measures": ["net_sales"]}, deps)
    charts, missing = charts_for(deps, ["q_day", "q_ch", "q_nope"])
    assert missing == ["q_nope"] and charts["q_ch"] == []  # a by-channel total is not a trend
    (chart, image), = charts["q_day"]
    assert image.cid == chart.cid and image.mime == "image/png" and image.content.startswith(b"\x89PNG")
