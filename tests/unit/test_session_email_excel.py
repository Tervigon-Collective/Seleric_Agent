"""The Excel attachment: readable, filterable, typed cells; sources and completeness recorded; text never a formula."""

from __future__ import annotations

import datetime as dt
from io import BytesIO

from openpyxl import load_workbook

from seleric_swarm.mail.excel import Column, Dataset, build_workbook


def _ds(**kw) -> Dataset:
    base = {
        "name": "Campaign spend",
        "columns": [Column("campaign", "Campaign"), Column("spend", "Spend (₹)", "currency"),
                 Column("ctr", "CTR", "percent"), Column("day", "Day", "date"), Column("orders", "Orders", "number")],
        "rows": [{"campaign": "Brand", "spend": 12345.6, "ctr": 0.0312, "day": "2026-10-27", "orders": 14},
              {"campaign": "Retarget", "spend": 800, "ctr": 0.05, "day": dt.date(2026, 10, 28), "orders": None}],
        "period": "2026-10-01 to 2026-10-27", "timezone": "Asia/Kolkata", "filters": "channel = meta",
        "source": "metrics_query q_ab12", "definitions": {"spend": "Ad spend, net of tax"},
    }
    return Dataset(**{**base, **kw})


def _open(ds: list[Dataset]):
    return load_workbook(BytesIO(build_workbook(ds)))


def test_data_sheet_has_frozen_header_filter_and_widths() -> None:
    ws = _open([_ds()])["Campaign spend"]
    assert [c.value for c in ws[1]] == ["Campaign", "Spend (₹)", "CTR", "Day", "Orders"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:E3"
    assert ws.column_dimensions["A"].width > 10 and ws.column_dimensions["B"].width >= len("Spend (₹)")


def test_cells_are_typed_and_formatted() -> None:
    ws = _open([_ds()])["Campaign spend"]
    assert ws["B2"].value == 12345.6 and "#,##0.00" in ws["B2"].number_format
    assert ws["C2"].value == 0.0312 and ws["C2"].number_format == "0.0%"
    assert isinstance(ws["D2"].value, dt.datetime) and ws["D2"].number_format == "yyyy-mm-dd"
    assert isinstance(ws["D3"].value, dt.datetime)  # a date object, not only an ISO string
    assert ws["E3"].value is None  # a missing value is blank, not 0


def test_text_that_looks_like_a_formula_stays_text() -> None:
    ds = _ds(rows=[{"campaign": '=HYPERLINK("http://evil","x")', "spend": 1}, {"campaign": "+1+1", "spend": 2},
                   {"campaign": "@SUM(A1)", "spend": 3}])
    ws = _open([ds])["Campaign spend"]
    for r in (2, 3, 4):
        assert ws.cell(r, 1).data_type == "s"
    assert ws["A2"].value == '=HYPERLINK("http://evil","x")'


def test_about_sheet_records_period_timezone_filters_source_definitions() -> None:
    about = _open([_ds()])["About"]
    text = "\n".join(str(c.value) for row in about.iter_rows() for c in row if c.value is not None)
    for expected in ("2026-10-01 to 2026-10-27", "Asia/Kolkata", "channel = meta", "metrics_query q_ab12",
                     "Ad spend, net of tax", "Complete", "2 rows"):
        assert expected in text


def test_partial_data_is_flagged_not_presented_as_complete() -> None:
    about = _open([_ds(total_rows=1534)])["About"]
    text = "\n".join(str(c.value) for row in about.iter_rows() for c in row if c.value is not None)
    assert "PARTIAL" in text and "2 of 1534" in text and "Complete" not in text


def test_sheet_names_are_valid_and_unique() -> None:
    wb = _open([_ds(name="Spend/ROAS: [Q4]? * very long name over the thirty one chars"), _ds(name="Spend/ROAS: [Q4]? * very long name over the thirty one chars")])
    names = wb.sheetnames
    assert len(set(names)) == 3 and all(len(n) <= 31 and not set("[]:*?/\\") & set(n) for n in names)
