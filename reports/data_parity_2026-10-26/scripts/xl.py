"""openpyxl helpers: styled tables with status colouring."""
import sys, os; sys.path.insert(0, os.path.dirname(__file__) + "/libs")
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
FILL = {"PASS": "C6EFCE", "WARN": "FFEB9C", "FAIL": "FFC7CE", "ERROR": "FF9999", "UNSTABLE": "BDD7EE", "INFO(partial)": "E7E6E6", "INFO": "E7E6E6",
        "NOT_RECOMPUTED": "D9D9D9", "CH_ERROR": "FF9999", "CUBE_ERROR": "FF9999", "High": "FFC7CE", "Medium": "FFEB9C", "Low": "E2EFDA", "Info": "E7E6E6"}
def fill_for(v):
    if not isinstance(v, str): return None
    if v in FILL: return FILL[v]
    if v.startswith("REFUSED"): return "F8CBAD"
    if v.startswith("PASS"): return "C6EFCE"
    if v.startswith("INFO"): return "E7E6E6"
    return None
HDR = PatternFill("solid", fgColor="1F3864"); HF = Font(bold=True, color="FFFFFF"); thin = Side(style="thin", color="BFBFBF")
def table(ws, headers, rows, widths=None, status_cols=(), start_row=1, numfmt="#,##0.00##", wrap_cols=(), freeze=True, autofilter=True):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h); c.fill = HDR; c.font = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            if isinstance(v, (list, dict, tuple)): v = str(v)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float): c.number_format = numfmt
            if (j - 1) in status_cols:
                f = fill_for(v)
                if f: c.fill = PatternFill("solid", fgColor=f); c.font = Font(bold=True)
            if (j - 1) in wrap_cols: c.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    if freeze: ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if autofilter and rows: ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
    return start_row + len(rows) + 1
def title(ws, text, row=1, size=14):
    c = ws.cell(row=row, column=1, value=text); c.font = Font(bold=True, size=size); return row + 1
